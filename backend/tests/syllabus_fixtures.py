from __future__ import annotations

import io

from docx import Document
from pypdf import PdfWriter


# a minimal one-page document with a real text object, so extraction has
# something to find without pulling in a pdf writing library
def text_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(out)


# what a scanned syllabus looks like to a text extractor
def blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def encrypted_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("a-password")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


# a syllabus states its weights in a table, so a fixture without one would not
# exercise the part of the reader that matters
def syllabus_docx(
    paragraphs: list[str], rows: list[tuple[str, str]] | None = None
) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if rows:
        table = document.add_table(rows=len(rows), cols=2)
        for index, (name, weight) in enumerate(rows):
            table.rows[index].cells[0].text = name
            table.rows[index].cells[1].text = weight
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def empty_docx() -> bytes:
    return syllabus_docx([])
